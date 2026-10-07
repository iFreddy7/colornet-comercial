"""Prepara a importação auditável. Não acessa o banco nem altera a planilha.
Uso: python script.py arquivo.xlsx catalogo.json clientes_localizacao.json saida.json
Preserva todas as linhas, inclusive itens iguais da mesma nota (não há número do item).
"""
import sys,json,hashlib,unicodedata,collections
from pathlib import Path
from datetime import datetime
from decimal import Decimal
import openpyxl

def norm(s):
    return ''.join(c for c in unicodedata.normalize('NFKD',str(s or '').strip().upper()) if not unicodedata.combining(c))

def preparar(arquivo,catalogo,localizacao):
    sha=hashlib.sha256(arquivo.read_bytes()).hexdigest()
    vendas={norm(v['nome']):v['id'] for v in catalogo['vendedores']}
    clientes=collections.defaultdict(list)
    produtos={p['codigo']:p['id'] for p in catalogo['produtos'] if p['codigo']}
    for c in localizacao:clientes[c['codigo'].strip()].append(c)
    w=openpyxl.load_workbook(arquivo,read_only=True,data_only=True)
    s=w.active
    assert list(next(s.iter_rows(values_only=True)))[:9]==['ANO','MESX','Filial','Numero','DT Emissao','VENDEDOR','CLIENTE','PRODUTO','QTDE'],'Cabeçalhos inesperados'
    resultado=[]
    for n,r in enumerate(s.iter_rows(min_row=2,values_only=True),2):
        if not any(x is not None for x in r):continue
        ano,mes=int(r[0]),int(r[1]);data=datetime.strptime(str(r[4]),'%d/%m/%Y')
        assert (ano,mes)==(data.year,data.month),f'Período divergente na linha {n}'
        cc,cn=map(str.strip,str(r[6]).split('-',1));pc,pn=map(str.strip,str(r[7]).split('-',1))
        vendedor=str(r[5] or '').strip();op=str(r[11] or '').strip();mun,uf=str(r[16] or '').strip(),str(r[17] or '').strip()
        assert not vendedor or norm(vendedor) in vendas,f'Vendedor não cadastrado: {vendedor}'
        qtd=Decimal(str(r[8]));assert qtd.is_finite() and qtd>=0 and qtd==qtd.quantize(Decimal('.0001'))
        candidatos=clientes[cc]
        if len(candidatos)>1:candidatos=[c for c in candidatos if norm(c['municipio'])==norm(mun) and norm(c['uf'])==norm(uf)]
        cliente_id=candidatos[0]['id'] if len(candidatos)==1 else None
        chave=cliente_id or '|'.join([cc,norm(mun),norm(uf),cn])
        excluido='Complemento fiscal ou de preço, sem quantidade vendida' if 'COMPLEM' in op else 'Venda de ativo imobilizado' if 'ATIVO IMOBILIZ' in op else None
        resultado.append(dict(arquivo_sha256=sha,linha=n,periodo=f'{ano}-{mes:02}-01',filial_codigo=str(r[2]),nota=str(r[3]),data_emissao=data.strftime('%Y-%m-%d'),vendedor_id=vendas.get(norm(vendedor)),vendedor_nome=vendedor or 'Sem vendedor informado',cliente_id=cliente_id,cliente_chave=chave,cliente_codigo=cc,cliente_nome=cn,municipio=mun,uf=uf,produto_id=produtos.get(pc),produto_codigo=pc,produto_nome=pn,quantidade=str(qtd),operacao=op,excluido_motivo=excluido))
    return {'lote':dict(arquivo_sha256=sha,arquivo_nome=arquivo.name,linhas_esperadas=len(resultado),quantidade_esperada=str(sum(Decimal(r['quantidade']) for r in resultado))), 'linhas':resultado}

if __name__=='__main__':
    arquivo,cat,loc,out=map(Path,sys.argv[1:])
    data=preparar(arquivo,json.loads(cat.read_text(encoding='utf8')),json.loads(loc.read_text(encoding='utf8')))
    out.write_text(json.dumps(data,ensure_ascii=False),encoding='utf8')
    print(json.dumps({'lote':data['lote'],'excluidas':sum(r['excluido_motivo'] is not None for r in data['linhas']),'sem_vendedor':sum(r['vendedor_id'] is None for r in data['linhas']),'sem_vinculo_cliente':sum(r['cliente_id'] is None for r in data['linhas'])}))
